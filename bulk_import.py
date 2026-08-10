"""Safe, append-only imports from updated AE Farmer Database workbooks."""
from __future__ import annotations

import json
from typing import Any, BinaryIO
from openpyxl import load_workbook
from db import set_setting, utc_now
from scripts.import_data import clean, json_value, normalize, normalize_age_group, normalize_phone, normalize_sex
from training import TRAINING_TOPICS, training_topic_status


class BulkImportError(ValueError):
    pass


def _identity_keys(core: dict[str, Any]) -> set[tuple[str, ...]]:
    keys = {(
        "profile", normalize(core.get("name")), normalize(core.get("sex")),
        normalize(core.get("group_name")), normalize(core.get("village")),
    )}
    phone = normalize_phone(core.get("phone"))
    if len(phone) >= 9:
        keys.add(("phone", phone))
    return keys


def _insert_record(connection, source_row: int, raw: dict[str, Any], core: dict[str, Any]) -> tuple[int, str]:
    timestamp = utc_now()
    cursor = connection.execute("INSERT INTO farmers(uid, created_at) VALUES(NULL, ?)", (timestamp,))
    farmer_id = cursor.lastrowid
    uid = f"ARF-{farmer_id:06d}"
    connection.execute("UPDATE farmers SET uid=? WHERE id=?", (uid, farmer_id))
    cursor = connection.execute(
        """INSERT INTO records(farmer_id, dataset, source_row, name, sex, age_value, age_group, phone,
           village, parish, subcounty, district, group_name, cbf_name, record_status, raw_data, created_at, updated_at)
           VALUES(?, 'training', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (farmer_id, source_row, core["name"], core["sex"], core["age_value"], core["age_group"], core["phone"],
         core["village"], core["parish"], core["subcounty"], core["district"], core["group_name"], core["cbf_name"],
         core["record_status"], json.dumps(raw, ensure_ascii=False, default=json_value), timestamp, timestamp))
    return cursor.lastrowid, uid


def append_farmer_database(connection, file: BinaryIO, filename: str = "") -> dict[str, int]:
    try:
        workbook = load_workbook(file, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise BulkImportError("The file is not a readable Excel workbook.") from exc
    try:
        if "Farmer Database" not in workbook.sheetnames:
            raise BulkImportError("The workbook must contain a 'Farmer Database' sheet.")
        sheet = workbook["Farmer Database"]
        headers = [clean(cell.value) for cell in next(sheet.iter_rows(min_row=2, max_row=2, max_col=88))]
        required = {"Name", "Sex", "Age group", "Phone Number", "Village", "GROUP NAME"}
        if not required.issubset(set(headers)):
            raise BulkImportError("This is not the expected Farmer Database format; required columns are missing.")
        cbf_map = {}
        if "CBF_Groups" in workbook.sheetnames:
            cbf_map = {normalize(group): clean(cbf) for cbf, group in
                       workbook["CBF_Groups"].iter_rows(min_row=2, max_col=2, values_only=True)
                       if clean(cbf) and clean(group)}
        existing_identities = set()
        for existing_row in connection.execute(
            "SELECT name, sex, phone, village, group_name FROM records WHERE dataset='training'"
        ).fetchall():
            existing_identities.update(_identity_keys(dict(existing_row)))
        next_source_row = connection.execute(
            "SELECT COALESCE(MAX(source_row), 2) + 1 FROM records WHERE dataset='training'"
        ).fetchone()[0]
        uploaded_identities = set()
        new_rows = []
        unchanged = duplicates = 0
        for row_number, row in enumerate(sheet.iter_rows(min_row=3, max_row=sheet.max_row, max_col=88, values_only=True), start=3):
            if not clean(row[0]):
                continue
            raw = {header: json_value(value) for header, value in zip(headers, row) if header}
            group_name = clean(raw.get("GROUP NAME"))
            core = {"name": clean(raw.get("Name")), "sex": normalize_sex(raw.get("Sex")),
                    "age_value": clean(raw.get("Age group")), "age_group": normalize_age_group(raw.get("Age group")),
                    "phone": clean(raw.get("Phone Number")), "village": clean(raw.get("Village")),
                    "parish": clean(raw.get("Parish")), "subcounty": clean(raw.get("Sub County")),
                    "district": clean(raw.get("DISTRICT")), "group_name": group_name,
                    "cbf_name": cbf_map.get(normalize(group_name), ""), "record_status": clean(raw.get("Status"))}
            if not core["name"]:
                raise BulkImportError(f"Row {row_number} has no beneficiary name.")
            identities = _identity_keys(core)
            if identities & existing_identities:
                unchanged += 1
                continue
            if identities & uploaded_identities:
                duplicates += 1
                continue
            uploaded_identities.update(identities)
            new_rows.append((raw, core))
        for raw, core in new_rows:
            assigned_source_row = next_source_row
            record_id, uid = _insert_record(connection, assigned_source_row, raw, core)
            next_source_row += 1
            core["audit"] = {
                "record_id": record_id, "uid": uid, "name": core["name"], "sex": core["sex"],
                "village": core["village"], "group_name": core["group_name"],
                "source_row": assigned_source_row,
            }
            for topic in TRAINING_TOPICS:
                status = training_topic_status(raw, topic)
                details = {key: value for key, value in status.items() if key not in {"status_code", "status_label", "training_received", "confirmed_trained", "followup_needed", "retraining_needed", "last_activity_date"}}
                connection.execute("""INSERT INTO topic_statuses(record_id, topic, status_code, status_label,
                    training_received, confirmed_trained, followup_needed, retraining_needed, last_activity_date, details)
                    VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (record_id, topic, status["status_code"], status["status_label"], status["training_received"],
                     status["confirmed_trained"], status["followup_needed"], status["retraining_needed"],
                     status.get("last_activity_date"), json.dumps(details, ensure_ascii=False) if details else None))
        if cbf_map:
            set_setting(connection, "cbf_group_map", cbf_map)
        set_setting(connection, "training_source", filename or "Farmer database upload")
        return {
            "added": len(new_rows), "existing": unchanged, "duplicates": duplicates,
            "beneficiaries": [core["audit"] for _raw, core in new_rows],
        }
    finally:
        workbook.close()
