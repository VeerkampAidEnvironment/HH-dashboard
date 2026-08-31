from __future__ import annotations

import json
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def iter_blocks(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == f"{{{W}}}p":
            yield Paragraph(child, parent)
        elif child.tag == f"{{{W}}}tbl":
            yield Table(child, parent)


def paragraph_record(paragraph: Paragraph) -> dict:
    p = paragraph._p
    visible = paragraph.text
    all_text = "".join(p.xpath(".//w:t/text()"))
    deleted = "".join(p.xpath(".//w:delText/text()"))
    inserted = "".join(p.xpath(".//w:ins//w:t/text()"))
    hyperlinks = []
    for h in p.xpath(".//w:hyperlink"):
        hyperlinks.append("".join(h.xpath(".//w:t/text()")))
    return {
        "kind": "paragraph",
        "style": paragraph.style.name if paragraph.style else None,
        "visible_text": visible,
        "all_current_text": all_text,
        "inserted_text": inserted,
        "deleted_text": deleted,
        "hyperlinks": hyperlinks,
    }


def table_record(table: Table) -> dict:
    rows = []
    for row in table.rows:
        rows.append(["\n".join(p.text for p in cell.paragraphs) for cell in row.cells])
    return {"kind": "table", "style": table.style.name if table.style else None, "rows": rows}


def xml_text(root) -> str:
    return "".join(root.xpath(".//w:t/text()", namespaces=NS))


def main() -> None:
    src = Path(sys.argv[1])
    out = Path(sys.argv[2])
    out.mkdir(parents=True, exist_ok=True)

    doc = Document(src)
    blocks = []
    for block in iter_blocks(doc):
        blocks.append(paragraph_record(block) if isinstance(block, Paragraph) else table_record(block))

    sections = []
    for idx, section in enumerate(doc.sections, start=1):
        sections.append({
            "section": idx,
            "page_width": section.page_width,
            "page_height": section.page_height,
            "top_margin": section.top_margin,
            "bottom_margin": section.bottom_margin,
            "left_margin": section.left_margin,
            "right_margin": section.right_margin,
            "header": "\n".join(p.text for p in section.header.paragraphs),
            "footer": "\n".join(p.text for p in section.footer.paragraphs),
        })

    package = {
        "source": str(src),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "sections": sections,
        "blocks": blocks,
    }

    media = []
    comments = []
    comments_extended = None
    tracked_change_counts = {"insertions": 0, "deletions": 0, "moves_from": 0, "moves_to": 0}
    with zipfile.ZipFile(src) as zf:
        names = set(zf.namelist())
        media_dir = out / "media"
        media_dir.mkdir(exist_ok=True)
        for name in sorted(n for n in names if n.startswith("word/media/") and not n.endswith("/")):
            target = media_dir / Path(name).name
            with zf.open(name) as r, target.open("wb") as w:
                shutil.copyfileobj(r, w)
            media.append({"package_path": name, "extracted_path": str(target), "size": target.stat().st_size})

        document_xml = etree.fromstring(zf.read("word/document.xml"))
        tracked_change_counts = {
            "insertions": len(document_xml.xpath(".//w:ins", namespaces=NS)),
            "deletions": len(document_xml.xpath(".//w:del", namespaces=NS)),
            "moves_from": len(document_xml.xpath(".//w:moveFrom", namespaces=NS)),
            "moves_to": len(document_xml.xpath(".//w:moveTo", namespaces=NS)),
        }

        if "word/comments.xml" in names:
            root = etree.fromstring(zf.read("word/comments.xml"))
            for c in root.xpath(".//w:comment", namespaces=NS):
                comments.append({
                    "id": c.get(f"{{{W}}}id"),
                    "author": c.get(f"{{{W}}}author"),
                    "date": c.get(f"{{{W}}}date"),
                    "text": xml_text(c),
                })
        if "word/commentsExtended.xml" in names:
            comments_extended = zf.read("word/commentsExtended.xml").decode("utf-8", errors="replace")

        package["package_parts"] = sorted(names)

    package["media"] = media
    package["comments"] = comments
    package["comments_extended_present"] = comments_extended is not None
    package["tracked_changes"] = tracked_change_counts

    (out / "review.json").write_text(json.dumps(package, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = []
    for idx, block in enumerate(blocks, start=1):
        if block["kind"] == "paragraph":
            text = block["visible_text"] or block["all_current_text"]
            if text.strip() or block["inserted_text"] or block["deleted_text"]:
                lines.append(f"[{idx:03d}] P ({block['style']}): {text}")
                if block["inserted_text"]:
                    lines.append(f"      INSERTED: {block['inserted_text']}")
                if block["deleted_text"]:
                    lines.append(f"      DELETED: {block['deleted_text']}")
        else:
            lines.append(f"[{idx:03d}] TABLE ({block['style']}):")
            for ridx, row in enumerate(block["rows"], start=1):
                lines.append(f"      R{ridx}: " + " || ".join(cell.replace("\n", " / ") for cell in row))
    (out / "content.txt").write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    main()
