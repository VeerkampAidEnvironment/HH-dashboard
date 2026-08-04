"""Import both supplied workbooks into the application database."""

from __future__ import annotations

import argparse
import json
import re
import sqlite3
import sys
import unicodedata
from datetime import date, datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from db import SCHEMA, set_setting  # noqa: E402
from training import TRAINING_TOPICS, care_module_status, training_topic_status  # noqa: E402


TRAINING_SOURCE = PROJECT_ROOT / "data" / "M_E_Dasboard_DONT_EDIT.xlsm"
CARE_SOURCE = PROJECT_ROOT / "data" / "Care groups monitoring DB.xlsx"
DEFAULT_DATABASE = PROJECT_ROOT / "instance" / "arfsa.db"

CARE_MODULE_RANGES = [
    ("Care groups - Module 1", "L", "O"),
    ("Care groups - Module 2", "P", "Z"),
    ("Care groups - Module 3", "AA", "AD"),
    ("Care groups - Module 4", "AE", "AG"),
    ("Care groups - Module 5", "AH", "AN"),
    ("School clubs - Nutrition comic book topics", "AR", "AV"),
    ("School clubs - Agronomic practices", "AW", "BB"),
    ("School clubs - Soil and water conservation", "BC", "BI"),
    ("School clubs - Theme 4", "BJ", "BO"),
    ("School clubs - Weeding, pest and disease control", "BP", "BT"),
    ("School clubs - Theme 7", "BU", "BW"),
]


def now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def json_value(value: Any):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKD", clean(value)).encode("ascii", "ignore").decode("ascii")
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text.lower()).split())


def normalize_phone(value: Any) -> str:
    digits = re.sub(r"\D", "", clean(value))
    if digits.startswith("256") and len(digits) >= 12:
        digits = "0" + digits[-9:]
    return digits


def normalize_sex(value: Any) -> str:
    text = normalize(value)
    if text in {"f", "female", "woman"}:
        return "F"
    if text in {"m", "male", "man"}:
        return "M"
    return clean(value).upper()[:20]


def normalize_age_group(value: Any, numeric_age: Any = None) -> str:
    text = normalize(value)
    mapping = {"y": "Youth", "youth": "Youth", "a": "Adult", "adult": "Adult", "e": "Elder", "elder": "Elder"}
    if text in mapping:
        return mapping[text]
    try:
        age = int(float(numeric_age if numeric_age not in (None, "") else value))
    except (TypeError, ValueError):
        return clean(value)
    if age <= 35:
        return "Youth"
    if age <= 59:
        return "Adult"
    return "Elder"


def create_farmer(connection: sqlite3.Connection) -> int:
    cursor = connection.execute("INSERT INTO farmers(uid, created_at) VALUES(NULL, ?)", (now(),))
    farmer_id = cursor.lastrowid
    connection.execute("UPDATE farmers SET uid=? WHERE id=?", (f"ARF-{farmer_id:06d}", farmer_id))
    return farmer_id


def insert_record(connection: sqlite3.Connection, dataset: str, source_row: int, raw: dict[str, Any],
                  core: dict[str, Any]) -> int:
    farmer_id = create_farmer(connection)
    timestamp = now()
    cursor = connection.execute(
        """INSERT INTO records(
            farmer_id, dataset, source_row, name, sex, age_value, age_group, phone,
            village, parish, subcounty, district, group_name, cbf_name, record_status,
            raw_data, created_at, updated_at
        ) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            farmer_id, dataset, source_row, core["name"], core.get("sex"), core.get("age_value"),
            core.get("age_group"), core.get("phone"), core.get("village"), core.get("parish"),
            core.get("subcounty"), core.get("district"), core.get("group_name"), core.get("cbf_name"),
            core.get("record_status"), json.dumps(raw, ensure_ascii=False, default=json_value), timestamp, timestamp,
        ),
    )
    return cursor.lastrowid


def insert_topic_status(connection: sqlite3.Connection, record_id: int, topic: str, status: dict[str, Any]) -> None:
    details = {key: value for key, value in status.items() if key not in {
        "status_code", "status_label", "training_received", "confirmed_trained",
        "followup_needed", "retraining_needed", "last_activity_date",
    }}
    connection.execute(
        """INSERT INTO topic_statuses(
            record_id, topic, status_code, status_label, training_received, confirmed_trained,
            followup_needed, retraining_needed, last_activity_date, details
        ) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            record_id, topic, status["status_code"], status["status_label"], status["training_received"],
            status["confirmed_trained"], status["followup_needed"], status["retraining_needed"],
            status.get("last_activity_date"), json.dumps(details, ensure_ascii=False) if details else None,
        ),
    )


def import_training(connection: sqlite3.Connection) -> int:
    workbook = load_workbook(TRAINING_SOURCE, read_only=True, data_only=True, keep_links=False)
    sheet = workbook["Farmer Database"]
    mapping_sheet = workbook["CBF_Groups"]
    cbf_map = {
        normalize(group): clean(cbf)
        for cbf, group in mapping_sheet.iter_rows(min_row=2, max_col=2, values_only=True)
        if clean(cbf) and clean(group)
    }
    set_setting(connection, "cbf_group_map", cbf_map)
    set_setting(connection, "training_topics", TRAINING_TOPICS)
    set_setting(connection, "training_source", TRAINING_SOURCE.name)

    headers = [clean(cell.value) for cell in next(sheet.iter_rows(min_row=2, max_row=2, max_col=88))]
    fields = []
    for header in headers:
        section = "Participant"
        if header.startswith("Training 1"):
            section = "Training cycle 1"
        elif header.startswith("Follow up 1"):
            section = "Follow-up cycle 1"
        elif header.startswith("Training 2"):
            section = "Training cycle 2"
        elif header.startswith("Follow up 2"):
            section = "Follow-up cycle 2"
        elif header.startswith("Training 3"):
            section = "Training cycle 3"
        elif header.startswith("Follow up 3"):
            section = "Follow-up cycle 3"
        field_type = "number" if "score -" in header.lower() or header == "Number of Birds Vaccinated" else "text"
        if header.startswith("Training ") or (header.startswith("Follow up ") and "score -" not in header.lower()):
            field_type = "date"
        fields.append({"key": header, "label": header, "section": section, "type": field_type})
    set_setting(connection, "training_schema", fields)

    imported = 0
    for row_number, row in enumerate(sheet.iter_rows(min_row=3, max_row=3000, max_col=88, values_only=True), start=3):
        if not clean(row[0]):
            continue
        raw = {header: json_value(value) for header, value in zip(headers, row) if header}
        group_name = clean(raw.get("GROUP NAME"))
        core = {
            "name": clean(raw.get("Name")),
            "sex": normalize_sex(raw.get("Sex")),
            "age_value": clean(raw.get("Age group")),
            "age_group": normalize_age_group(raw.get("Age group")),
            "phone": clean(raw.get("Phone Number")),
            "village": clean(raw.get("Village")),
            "parish": clean(raw.get("Parish")),
            "subcounty": clean(raw.get("Sub County")),
            "district": clean(raw.get("DISTRICT")),
            "group_name": group_name,
            "cbf_name": cbf_map.get(normalize(group_name), ""),
            "record_status": clean(raw.get("Status")),
        }
        record_id = insert_record(connection, "training", row_number, raw, core)
        for topic in TRAINING_TOPICS:
            insert_topic_status(connection, record_id, topic, training_topic_status(raw, topic))
        imported += 1
    workbook.close()
    return imported


def care_field_schema(sheet) -> tuple[list[str], list[dict[str, Any]], dict[str, list[str]]]:
    headers: list[str] = []
    fields: list[dict[str, Any]] = []
    module_fields: dict[str, list[str]] = {name: [] for name, _, _ in CARE_MODULE_RANGES}
    row2 = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2, max_col=76))]
    row3 = [cell.value for cell in next(sheet.iter_rows(min_row=3, max_row=3, max_col=76))]
    basic_labels = {
        1: "No.", 2: "Project Participant", 3: "Sex", 4: "Age", 5: "Category",
        6: "Name of group/school", 7: "Village", 8: "Parish", 9: "Subcounty", 10: "District",
        11: "Contact", 41: "Status", 42: "Month dropped", 43: "Replacement",
    }
    for index in range(1, 77):
        module_name = None
        for name, start, end in CARE_MODULE_RANGES:
            from openpyxl.utils.cell import column_index_from_string
            if column_index_from_string(start) <= index <= column_index_from_string(end):
                module_name = name
                break
        if index in basic_labels:
            key = basic_labels[index]
            section = "Participant"
            field_type = "number" if index in {1, 4} else "text"
        elif module_name:
            session = clean(row3[index - 1]) or str(index)
            key = f"{module_name} - Session {session}"
            section = module_name
            field_type = "attendance"
            module_fields[module_name].append(key)
        else:
            label = clean(row2[index - 1]) or f"Column {get_column_letter(index)}"
            key = label
            section = "Other"
            field_type = "text"
        headers.append(key)
        fields.append({"key": key, "label": key, "section": section, "type": field_type})
    return headers, fields, module_fields


def import_care(connection: sqlite3.Connection) -> int:
    workbook = load_workbook(CARE_SOURCE, read_only=True, data_only=True, keep_links=False)
    sheet = workbook["Sheet1"]
    headers, fields, module_fields = care_field_schema(sheet)
    set_setting(connection, "care_schema", fields)
    set_setting(connection, "care_modules", module_fields)
    set_setting(connection, "care_source", CARE_SOURCE.name)
    imported = 0
    for row_number, row in enumerate(sheet.iter_rows(min_row=4, max_row=sheet.max_row, max_col=76, values_only=True), start=4):
        if not clean(row[1]):
            continue
        raw = {header: json_value(value) for header, value in zip(headers, row)}
        core = {
            "name": clean(raw.get("Project Participant")),
            "sex": normalize_sex(raw.get("Sex")),
            "age_value": clean(raw.get("Age")),
            "age_group": normalize_age_group("", raw.get("Age")),
            "phone": clean(raw.get("Contact")),
            "village": clean(raw.get("Village")),
            "parish": clean(raw.get("Parish")),
            "subcounty": clean(raw.get("Subcounty")),
            "district": clean(raw.get("District")),
            "group_name": clean(raw.get("Name of group/school")),
            "cbf_name": "",
            "record_status": clean(raw.get("Status")),
        }
        record_id = insert_record(connection, "care", row_number, raw, core)
        for module_name, field_keys in module_fields.items():
            insert_topic_status(connection, record_id, module_name, care_module_status(raw, field_keys))
        imported += 1
    workbook.close()
    return imported


def candidate_match(training: sqlite3.Row, care: sqlite3.Row) -> tuple[int, list[str], str] | None:
    t_name, c_name = normalize(training["name"]), normalize(care["name"])
    if not t_name or not c_name:
        return None
    t_phone, c_phone = normalize_phone(training["phone"]), normalize_phone(care["phone"])
    if t_phone and c_phone and t_phone != c_phone:
        return None
    name_similarity = SequenceMatcher(None, t_name, c_name).ratio()
    same_name = t_name == c_name
    same_phone = len(t_phone) >= 9 and t_phone == c_phone
    same_sex = bool(training["sex"] and training["sex"] == care["sex"])
    locations = [
        (normalize(training["village"]), normalize(care["village"]), "village"),
        (normalize(training["parish"]), normalize(care["parish"]), "parish"),
        (normalize(training["subcounty"]), normalize(care["subcounty"]), "subcounty"),
    ]
    same_locations = [label for left, right, label in locations if left and left == right]
    if same_name and same_phone:
        return 100, ["exact name", "exact phone"], "pending"
    if same_name and same_sex and same_locations:
        return 96, ["exact name", "same sex", f"same {same_locations[0]}"], "pending"
    if same_phone and name_similarity >= 0.78:
        return 92, ["exact phone", f"similar name ({name_similarity:.0%})"], "pending"
    if same_name and same_sex:
        return 84, ["exact name", "same sex"], "pending"
    return None


def build_matches(connection: sqlite3.Connection) -> tuple[int, int]:
    training_rows = connection.execute("SELECT * FROM records WHERE dataset='training'").fetchall()
    care_rows = connection.execute("SELECT * FROM records WHERE dataset='care'").fetchall()
    care_by_name: dict[str, list[sqlite3.Row]] = {}
    care_by_phone: dict[str, list[sqlite3.Row]] = {}
    for care in care_rows:
        care_by_name.setdefault(normalize(care["name"]), []).append(care)
        phone = normalize_phone(care["phone"])
        if len(phone) >= 9:
            care_by_phone.setdefault(phone, []).append(care)
    confirmed = pending = 0
    used_care: set[int] = set()
    for training in training_rows:
        candidates = []
        candidate_pool: dict[int, sqlite3.Row] = {}
        for care in care_by_name.get(normalize(training["name"]), []):
            candidate_pool[care["id"]] = care
        phone = normalize_phone(training["phone"])
        if len(phone) >= 9:
            for care in care_by_phone.get(phone, []):
                candidate_pool[care["id"]] = care
        for care in candidate_pool.values():
            if care["id"] in used_care:
                continue
            candidate = candidate_match(training, care)
            if candidate:
                candidates.append((candidate[0], care, candidate[1], candidate[2]))
        if not candidates:
            continue
        candidates.sort(key=lambda item: item[0], reverse=True)
        score, care, reasons, status = candidates[0]
        if len(candidates) > 1 and candidates[1][0] == score:
            status = "pending"
            reasons.append("multiple equally strong candidates")
        connection.execute(
            "INSERT INTO matches(training_record_id, care_record_id, confidence, reasons, status, reviewed_at) "
            "VALUES(?, ?, ?, ?, ?, ?)",
            (
                training["id"], care["id"], score, json.dumps(reasons), status,
                now() if status == "confirmed" else None,
            ),
        )
        if status == "confirmed":
            old_farmer_id = care["farmer_id"]
            connection.execute("UPDATE records SET farmer_id=? WHERE id=?", (training["farmer_id"], care["id"]))
            connection.execute(
                "DELETE FROM farmers WHERE id=? AND NOT EXISTS(SELECT 1 FROM records WHERE farmer_id=?)",
                (old_farmer_id, old_farmer_id),
            )
            used_care.add(care["id"])
            confirmed += 1
        else:
            pending += 1
    return confirmed, pending


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE)
    parser.add_argument("--reset", action="store_true", help="Replace imported database content")
    args = parser.parse_args()
    args.database.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(args.database)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.executescript(SCHEMA)
    if args.reset:
        connection.executescript(
            "DELETE FROM audit_log; DELETE FROM matches; DELETE FROM topic_statuses; "
            "DELETE FROM records; DELETE FROM farmers; DELETE FROM settings; "
            "DELETE FROM sqlite_sequence WHERE name IN "
            "('audit_log','matches','topic_statuses','records','farmers');"
        )
    elif connection.execute("SELECT COUNT(*) FROM records").fetchone()[0]:
        raise SystemExit("Database already contains records. Use --reset only when replacement is intended.")
    training_count = import_training(connection)
    care_count = import_care(connection)
    confirmed, pending = build_matches(connection)
    set_setting(connection, "imported_at", now())
    set_setting(connection, "status_calculated_on", date.today().isoformat())
    set_setting(connection, "import_summary", {
        "training_records": training_count,
        "care_records": care_count,
        "confirmed_matches": confirmed,
        "pending_matches": pending,
    })
    connection.commit()
    unique_farmers = connection.execute("SELECT COUNT(*) FROM farmers").fetchone()[0]
    connection.close()
    print(json.dumps({
        "database": str(args.database.resolve()),
        "training_records": training_count,
        "care_records": care_count,
        "confirmed_matches": confirmed,
        "pending_matches": pending,
        "unique_farmers": unique_farmers,
    }, indent=2))


if __name__ == "__main__":
    main()
