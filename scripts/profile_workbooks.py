"""Produce a compact, read-only structural profile of the project workbooks."""

from __future__ import annotations

import collections
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


KEYWORDS = re.compile(
    r"train|follow|status|score|cbf|facilit|farmer|group|visit|month|day|threshold",
    re.IGNORECASE,
)


def scalar(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def normalized_formula(formula: str) -> str:
    return re.sub(r"(?<![A-Z_])\$?[A-Z]{1,3}\$?\d+", "<CELL>", formula.upper())[:500]


def profile(path: Path) -> dict:
    workbook = load_workbook(
        path,
        read_only=False,
        data_only=False,
        keep_vba=path.suffix.lower() == ".xlsm",
        keep_links=True,
    )
    result = {
        "file": path.name,
        "sheet_count": len(workbook.sheetnames),
        "defined_names": [],
        "sheets": [],
    }

    for item in workbook.defined_names.values():
        result["defined_names"].append(
            {
                "name": item.name,
                "value": scalar(item.attr_text),
                "hidden": bool(item.hidden),
                "local_sheet_id": item.localSheetId,
            }
        )

    for sheet in workbook.worksheets:
        formulas = []
        keyword_cells = []
        row_summaries = []
        formula_patterns = collections.Counter()
        nonempty_count = 0

        for row in sheet.iter_rows():
            row_items = []
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                nonempty_count += 1
                text = str(value)
                if len(row_items) < 50:
                    row_items.append({"cell": cell.coordinate, "value": scalar(value)})
                if text.startswith("="):
                    formula_patterns[normalized_formula(text)] += 1
                    if len(formulas) < 80:
                        formulas.append({"cell": cell.coordinate, "formula": text[:1000]})
                if KEYWORDS.search(text) and len(keyword_cells) < 250:
                    keyword_cells.append({"cell": cell.coordinate, "value": text[:1000]})
            if row_items and len(row_summaries) < 25:
                row_summaries.append(row_items)

        tables = []
        for table in sheet.tables.values():
            tables.append(
                {
                    "name": table.name,
                    "display_name": table.displayName,
                    "ref": table.ref,
                    "columns": [column.name for column in table.tableColumns],
                }
            )

        validations = []
        if sheet.data_validations:
            for validation in sheet.data_validations.dataValidation:
                validations.append(
                    {
                        "ranges": str(validation.sqref),
                        "type": validation.type,
                        "formula1": scalar(validation.formula1),
                        "formula2": scalar(validation.formula2),
                    }
                )

        conditional_formats = []
        for conditional_range in sheet.conditional_formatting:
            for rule in conditional_range.rules:
                conditional_formats.append(
                    {
                        "range": str(conditional_range.sqref),
                        "type": rule.type,
                        "operator": rule.operator,
                        "formula": [scalar(value) for value in (rule.formula or [])],
                    }
                )

        result["sheets"].append(
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "dimension": sheet.calculate_dimension(),
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "nonempty_cells": nonempty_count,
                "freeze_panes": scalar(sheet.freeze_panes),
                "auto_filter": scalar(sheet.auto_filter.ref),
                "merged_ranges": [str(item) for item in list(sheet.merged_cells.ranges)[:100]],
                "table_count": len(tables),
                "tables": tables,
                "chart_count": len(sheet._charts),
                "image_count": len(sheet._images),
                "formula_count": sum(formula_patterns.values()),
                "formula_patterns": [
                    {"count": count, "formula": formula}
                    for formula, count in formula_patterns.most_common(30)
                ],
                "formula_samples": formulas,
                "keyword_cells": keyword_cells,
                "top_nonempty_rows": row_summaries,
                "validations": validations[:100],
                "conditional_formats": conditional_formats[:200],
            }
        )

    workbook.close()
    return result


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("Usage: profile_workbooks.py FILE [FILE ...]")
    profiles = [profile(Path(argument)) for argument in sys.argv[1:]]
    print(json.dumps(profiles, indent=2, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
